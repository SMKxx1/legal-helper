"""Template-studio page flows (P5 wave B) — upload → view → tokenize → checklist → publish/rollback,
undo/redo, find-and-map, test-drive, and the stale-view refusal — end-to-end against a seeded draft.

Mounts the REAL app (``client`` fixture: middleware, auth deps, error envelope) on a throwaway DB.
The ``login`` fixture logs an admin in and sets the CSRF header, so state-changing /admin POSTs pass
both the session gate and the double-submit check.
"""

from __future__ import annotations

from io import BytesIO

from conftest_admin import (
    add_version,
    all_tokens_docx,
    bare_docx,
    seed_catalog,
    seed_draft,
    source_docx,
    span_of,
    template_id,
    tokens_docx,
)
from docx import Document

from app.models_v2 import TemplateVersion
from app.studio.checklist import scan_token_names
from app.studio.docview import extract_view
from app.studio.findmap import detect_placeholders


def _admin(client, seed_user, login):
    seed_user("admin", role="admin")
    assert login("admin").status_code == 200


# --------------------------------------------------------------------------- #
# Templates list + upload
# --------------------------------------------------------------------------- #
def test_templates_list_renders_slots(client, db, seed_user, login):
    seed_catalog(db)
    _admin(client, seed_user, login)
    resp = client.get("/admin/templates")
    assert resp.status_code == 200
    # All 8 templates listed, each with an empty + tokenised variant + an upload form.
    assert resp.text.count('class="tpl-upload"') == 16  # 8 templates × 2 variants
    assert "US" in resp.text and "ServiceProvider" in resp.text


def test_upload_creates_draft_and_opens_studio(client, db, seed_user, login):
    seed_catalog(db)
    _admin(client, seed_user, login)
    tid = template_id(db)
    resp = client.post(
        f"/admin/templates/{tid}/tokenised/upload",
        files={
            "file": (
                "nda.docx",
                source_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 201
    body = resp.json()
    assert body["ok"] is True
    version_id = body["version_id"]
    assert body["studio_url"] == f"/admin/studio/{version_id}"

    # The new draft is not current, and the studio page renders the uploaded document.
    version = db.get(TemplateVersion, version_id)
    assert version.is_current is False
    assert (
        version.created_by == "admin"
    )  # attribution (P6): the uploading admin user id
    page = client.get(body["studio_url"])
    assert page.status_code == 200
    assert "ACME CORPORATION" in page.text
    assert 'class="stu-seg"' in page.text or "stu-seg" in page.text


def test_upload_rejects_non_docx(client, db, seed_user, login):
    seed_catalog(db)
    _admin(client, seed_user, login)
    tid = template_id(db)
    resp = client.post(
        f"/admin/templates/{tid}/tokenised/upload",
        files={"file": ("nda.txt", b"not a docx", "text/plain")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "not_docx"


def test_studio_page_needs_upload_when_no_blob(client, db, seed_user, login):
    seed_catalog(db)
    _admin(client, seed_user, login)
    # The seeded tokenised v1 has is_current but no blob loaded.
    tid = template_id(db)
    version = db.execute(
        TemplateVersion.__table__.select().where(
            (TemplateVersion.template_id == tid)
            & (TemplateVersion.variant_code == "tokenised")
        )
    ).first()
    resp = client.get(f"/admin/studio/{version.id}")
    assert resp.status_code == 200
    assert "No document yet" in resp.text


# --------------------------------------------------------------------------- #
# Studio editor render + tokenize
# --------------------------------------------------------------------------- #
def test_studio_page_renders_palette_and_checklist(client, db, seed_user, login):
    _, version_id = seed_draft(db, source_docx())
    _admin(client, seed_user, login)
    resp = client.get(f"/admin/studio/{version_id}")
    assert resp.status_code == 200
    assert "stu-palette" in resp.text
    assert "stu-chip" in resp.text  # registry palette rendered
    assert 'data-locator="body/p:0"' in resp.text
    assert 'class="stu-run"' in resp.text  # doc text rendered as data-plen part nodes
    assert "counterparty_name" in resp.text  # a seed token in the palette
    # doc view carries the view hash for tokenize ops
    view = extract_view(source_docx())
    assert view.content_hash in resp.text


def test_tokenize_happy_path(client, db, seed_user, login):
    _, version_id = seed_draft(db, source_docx())
    _admin(client, seed_user, login)
    start, end, view_hash = span_of(source_docx(), "body/p:0", "ACME CORPORATION")
    resp = client.post(
        f"/admin/studio/{version_id}/tokenize",
        json={
            "locator": "body/p:0",
            "start": start,
            "end": end,
            "token": "counterparty_name",
            "view_hash": view_hash,
        },
    )
    assert resp.status_code == 200
    st = resp.json()
    # The token renders as an ATOMIC chip showing the friendly registry label — the raw
    # {{counterparty_name}} braces never appear as visible document text (only as the chip's
    # tooltip), and the chip carries the underlying length for the client's offset arithmetic.
    assert 'data-token="counterparty_name"' in st["doc_html"]
    assert (
        ">Counterparty Name</span>" in st["doc_html"]
    )  # seeded registry label, not snake_case
    assert ">{{counterparty_name}}</span>" not in st["doc_html"]
    assert f'data-plen="{len("{{counterparty_name}}")}"' in st["doc_html"]
    assert "counterparty_name" not in st["missing_required"]  # now satisfied
    assert st["can_undo"] is True
    # The draft blob really changed (persisted through the oplog).
    version = db.get(TemplateVersion, version_id)
    db.expire(version)
    from app.models_v2 import DocumentBlob

    blob = db.get(DocumentBlob, db.get(TemplateVersion, version_id).blob_id)
    assert "counterparty_name" in scan_token_names(blob.bytes)


def test_tokenize_stale_view_is_409(client, db, seed_user, login):
    _, version_id = seed_draft(db, source_docx())
    _admin(client, seed_user, login)
    start, end, _ = span_of(source_docx(), "body/p:0", "ACME CORPORATION")
    resp = client.post(
        f"/admin/studio/{version_id}/tokenize",
        json={
            "locator": "body/p:0",
            "start": start,
            "end": end,
            "token": "counterparty_name",
            "view_hash": "0" * 64,  # not the current content hash
        },
    )
    assert resp.status_code == 409
    body = resp.json()
    assert body["error"]["code"] == "studio_stale_view"
    assert "expected_hash" in body["error"]["details"]
    assert "actual_hash" in body["error"]["details"]


def test_tokenize_refusal_passthrough_empty_span(client, db, seed_user, login):
    _, version_id = seed_draft(db, source_docx())
    _admin(client, seed_user, login)
    _, _, view_hash = span_of(source_docx(), "body/p:0", "ACME CORPORATION")
    # A zero-width selection → the studio empty-span refusal, surfaced in the envelope.
    resp = client.post(
        f"/admin/studio/{version_id}/tokenize",
        json={
            "locator": "body/p:0",
            "start": 5,
            "end": 5,
            "token": "counterparty_name",
            "view_hash": view_hash,
        },
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "studio_empty_span"


def test_tokenize_cross_paragraph_refusal(client, db, seed_user, login):
    _, version_id = seed_draft(db, source_docx())
    _admin(client, seed_user, login)
    start, end, view_hash = span_of(source_docx(), "body/p:0", "ACME CORPORATION")
    resp = client.post(
        f"/admin/studio/{version_id}/tokenize",
        json={
            "locator": "body/p:0",
            "end_locator": "body/p:1",
            "start": start,
            "end": end,
            "token": "counterparty_name",
            "view_hash": view_hash,
        },
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "studio_cross_paragraph"


# --------------------------------------------------------------------------- #
# Undo / redo via the routes
# --------------------------------------------------------------------------- #
def test_undo_then_redo(client, db, seed_user, login):
    _, version_id = seed_draft(db, source_docx())
    _admin(client, seed_user, login)
    start, end, view_hash = span_of(source_docx(), "body/p:0", "ACME CORPORATION")
    client.post(
        f"/admin/studio/{version_id}/tokenize",
        json={
            "locator": "body/p:0",
            "start": start,
            "end": end,
            "token": "counterparty_name",
            "view_hash": view_hash,
        },
    )
    undo = client.post(f"/admin/studio/{version_id}/undo", json={})
    assert undo.status_code == 200
    assert (
        "counterparty_name" not in undo.json()["doc_html"]
    )  # chip fully gone after undo
    assert undo.json()["can_redo"] is True

    redo = client.post(f"/admin/studio/{version_id}/redo", json={})
    assert redo.status_code == 200
    assert 'data-token="counterparty_name"' in redo.json()["doc_html"]  # chip is back


def test_undo_with_nothing_is_409(client, db, seed_user, login):
    _, version_id = seed_draft(db, source_docx())
    _admin(client, seed_user, login)
    resp = client.post(f"/admin/studio/{version_id}/undo", json={})
    assert resp.status_code == 409
    assert resp.json()["error"]["code"] == "studio_nothing_to_undo"


# --------------------------------------------------------------------------- #
# Find-and-map assistant
# --------------------------------------------------------------------------- #
def test_findmap_accept_maps_placeholder(client, db, seed_user, login):
    _, version_id = seed_draft(db, source_docx())
    _admin(client, seed_user, login)
    view = extract_view(source_docx())
    cands = detect_placeholders(view, [("effective_date", "Effective date")])
    cand = next(c for c in cands if c.suggested_token == "effective_date")
    resp = client.post(
        f"/admin/studio/{version_id}/map",
        json={
            "view_hash": view.content_hash,
            "mappings": [
                {
                    "locator": cand.locator,
                    "start": cand.start,
                    "end": cand.end,
                    "token_name": "effective_date",
                }
            ],
        },
    )
    assert resp.status_code == 200
    assert 'data-token="effective_date"' in resp.json()["doc_html"]


# --------------------------------------------------------------------------- #
# Publish gating + drift
# --------------------------------------------------------------------------- #
def test_publish_blocked_when_required_missing(client, db, seed_user, login):
    _, version_id = seed_draft(db, bare_docx())
    _admin(client, seed_user, login)
    resp = client.post(f"/admin/studio/{version_id}/publish", json={})
    assert resp.status_code == 409
    body = resp.json()
    assert body["error"]["code"] == "publish_blocked"
    assert body["error"]["details"]["missing_required"]  # non-empty


def test_publish_success_flips_current_and_emits_drift(client, db, seed_user, login):
    tid, version_id = seed_draft(db, all_tokens_docx())
    _admin(client, seed_user, login)
    resp = client.post(f"/admin/studio/{version_id}/publish", json={})
    assert resp.status_code == 200
    body = resp.json()
    assert body["ok"] is True
    assert body["added_tokens"]  # the newly-published token set

    # is_current flipped to the published draft; the superseded drafts (the seed v1) are DELETED —
    # once a version is made current only the live version remains for its (template, variant).
    db.expire_all()
    published = db.get(TemplateVersion, version_id)
    assert published.is_current is True
    remaining = (
        db.query(TemplateVersion)
        .filter(
            TemplateVersion.template_id == tid,
            TemplateVersion.variant_code == "tokenised",
        )
        .all()
    )
    assert [v.id for v in remaining] == [version_id]


def test_rollback_repoints_current_and_emits_drift(client, db, seed_user, login):
    from sqlalchemy import update

    seed_catalog(db)
    tid = template_id(db)
    # Single clean current: clear the seed's current, then a current Y and an older non-current X.
    db.execute(
        update(TemplateVersion)
        .where(
            TemplateVersion.template_id == tid,
            TemplateVersion.variant_code == "tokenised",
        )
        .values(is_current=False)
    )
    db.commit()
    current_y = add_version(
        db, tid, "tokenised", tokens_docx("counterparty_name"), is_current=True
    )
    older_x = add_version(
        db, tid, "tokenised", tokens_docx("effective_date"), is_current=False
    )
    _admin(client, seed_user, login)

    resp = client.post(f"/admin/templates/versions/{older_x}/rollback", json={})
    assert resp.status_code == 200
    body = resp.json()
    assert body["ok"] is True
    assert body["added_tokens"] == ["effective_date"]
    assert body["removed_tokens"] == ["counterparty_name"]

    # Making older_x current deletes the superseded version (current_y) — only the live one remains.
    db.expire_all()
    assert db.get(TemplateVersion, current_y) is None
    rolled = db.get(TemplateVersion, older_x)
    assert rolled is not None and rolled.is_current is True
    remaining = (
        db.query(TemplateVersion)
        .filter(
            TemplateVersion.template_id == tid,
            TemplateVersion.variant_code == "tokenised",
        )
        .all()
    )
    assert [v.id for v in remaining] == [older_x]


# --------------------------------------------------------------------------- #
# Test-drive
# --------------------------------------------------------------------------- #
def test_test_drive_returns_filled_docx(client, db, seed_user, login):
    _, version_id = seed_draft(
        db, tokens_docx("counterparty_name", "effective_date", "notice_email")
    )
    _admin(client, seed_user, login)
    resp = client.get(f"/admin/studio/{version_id}/test-drive")
    assert resp.status_code == 200
    assert "wordprocessingml" in resp.headers["content-type"]
    assert "attachment" in resp.headers.get("content-disposition", "")
    # A real, readable .docx whose tokens are all filled (dummy values), none left as {{token}}.
    filled = resp.content
    Document(BytesIO(filled))  # parses without error
    assert scan_token_names(filled) == []
