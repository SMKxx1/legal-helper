"""Envelope interactivity chains (PLAN §3.9, reference §3.7) — send_docusign / modal / confirm / use-doc.

Drives :func:`app.bot.interactivity.dispatch_interaction` with hand-built Slack bodies + a captured Slack
sink (a stub ``WebClient``) and INJECTED envelope collaborators (a fake DocuSign sender, a capturing
modal opener, a fake Slack fetcher) — zero network. Covers: the ``send_docusign`` modal open, the
``nda_docusign`` modal golden parse + validation, the *Confirm & send* / *Cancel* clicks (with the
requester-mapping ``nda_envelopes`` row assertions, PLAN §3.10), the confirm-before-send invariant (no
send on the modal/collect steps), idempotent double-click, capability-off degradation, failed-attempt
persistence, and the ``env_use_doc`` / ``decline_doc`` thread-doc chain.
"""

from __future__ import annotations

import base64
import json
from io import BytesIO
from typing import Any

from docx import Document

from app.bot.blockkit import (
    decline_doc_value,
    env_confirm_value,
    env_use_doc_value,
    send_docusign_value,
)
from app.bot.channels.replies import ReplyService, SlackReplySink
from app.bot.intents.envelope import (
    ALREADY_SENT_TEXT,
    CANCEL_TEXT,
    DECLINE_TEXT,
    DOCUSIGN_UNAVAILABLE_TEXT,
    EXPIRED_STATE_TEXT,
    MODAL_INVALID_TEXT,
    EnvelopeDeps,
    register_envelope,
)
from app.bot.interactivity import (
    InteractivityDeps,
    default_interactivity_registry,
    dispatch_interaction,
)
from app.bot.models import BotCorrelation
from app.config import Settings
from app.integrations.docusign import DocuSignTerminalError, EnvelopeResult
from app.integrations.models import NdaEnvelope

pytest_plugins = ("conftest_bot",)


# --------------------------------------------------------------------------- #
# .docx fixtures
# --------------------------------------------------------------------------- #
def _docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    for para in paragraphs:
        doc.add_paragraph(para)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


CLEAN_DOCX = _docx(["A finished NDA between Amperesand and Acme."])
TOKENISED_DOCX = _docx(["Between {{amperesand_signer_name}} and Acme."])


# --------------------------------------------------------------------------- #
# Test doubles
# --------------------------------------------------------------------------- #
class StubWebClient:
    def __init__(self) -> None:
        self.messages: list[dict[str, Any]] = []

    def chat_postMessage(self, **kw: Any) -> dict[str, Any]:
        self.messages.append(kw)
        return {"ts": "resp.ts"}


class FakeSender:
    """A fake ``create_and_send_envelope`` — records kwargs, returns a result or raises."""

    def __init__(
        self, *, result: EnvelopeResult | None = None, exc: Exception | None = None
    ) -> None:
        self.calls: list[dict[str, Any]] = []
        self._result = result
        self._exc = exc

    def __call__(self, **kw: Any) -> EnvelopeResult:
        self.calls.append(kw)
        if self._exc is not None:
            raise self._exc
        assert self._result is not None
        return self._result


def _deps(
    bot_session_factory: Any, *, settings: Settings | None = None
) -> tuple[InteractivityDeps, StubWebClient]:
    stub = StubWebClient()
    sink = SlackReplySink(stub)
    service = ReplyService([sink])
    settings = settings or Settings(  # type: ignore[call-arg]
        _env_file=None, nda_bot_from_email="nda-bot@example.com"
    )
    deps = InteractivityDeps(
        session_factory=bot_session_factory,
        service=service,
        post_blocks=sink.post_blocks,
        settings=settings,
    )
    return deps, stub


def _run(
    bot_session_factory: Any,
    body: dict[str, Any],
    env_deps: EnvelopeDeps,
    *,
    settings: Settings | None = None,
) -> StubWebClient:
    deps, stub = _deps(bot_session_factory, settings=settings)
    registry = default_interactivity_registry()
    register_envelope(
        registry, deps=env_deps
    )  # last-wins over the default envelope registration
    dispatch_interaction(body, registry=registry, deps=deps)
    return stub


# --------------------------------------------------------------------------- #
# State seeding + assertions
# --------------------------------------------------------------------------- #
def _state(**over: Any) -> dict[str, Any]:
    base = {
        "file_name": "nda.docx",
        "channel": "slack",
        "slack_channel": "C1",
        "slack_thread_ts": "T1",
        "email_message_id": "",
        "requested_by": "U1",
        "signer_emails": ["amp@a.com", "cp@b.com"],
        "cc_emails": [],
        "routing": "amp_first",
        "cc_timing": "after",
        "doc_b64": base64.b64encode(CLEAN_DOCX).decode("ascii"),
    }
    base.update(over)
    return base


def _seed(bot_session_factory: Any, ref: str, payload: dict[str, Any]) -> None:
    with bot_session_factory() as s, s.begin():
        s.add(BotCorrelation(key=ref, kind="env_confirm", payload_json=payload))


def _load(bot_session_factory: Any, ref: str) -> dict[str, Any] | None:
    with bot_session_factory() as s:
        row = s.query(BotCorrelation).filter(BotCorrelation.key == ref).one_or_none()
        return dict(row.payload_json or {}) if row else None


def _envelopes(bot_session_factory: Any) -> list[NdaEnvelope]:
    with bot_session_factory() as s:
        return list(s.query(NdaEnvelope).all())


def _buttons(blocks: Any) -> dict[str, str]:
    out: dict[str, str] = {}
    for block in blocks or ():
        if block.get("type") == "actions":
            for el in block.get("elements", []):
                if el.get("type") == "button":
                    out[el["action_id"]] = el.get("value", "")
    return out


# --------------------------------------------------------------------------- #
# Body builders
# --------------------------------------------------------------------------- #
def _actions_body(
    action_id: str,
    value: str,
    *,
    channel: str = "C1",
    thread_ts: str = "T1",
    message_ts: str = "M1",
    clicker: str = "U1",
    trigger_id: str = "trig",
) -> dict[str, Any]:
    return {
        "type": "block_actions",
        "user": {"id": clicker},
        "channel": {"id": channel},
        "container": {
            "type": "message",
            "message_ts": message_ts,
            "thread_ts": thread_ts,
            "channel_id": channel,
        },
        "message": {"ts": message_ts, "thread_ts": thread_ts},
        "trigger_id": trigger_id,
        "actions": [{"action_id": action_id, "type": "button", "value": value}],
        "state": {"values": {}},
    }


def _modal_body(
    ref: str,
    *,
    amp: str = "amp@a.com",
    cp: str = "cp@b.com",
    order: str = "cp_first",
    cc: str = "x@y.com, z@w.com; q@r.com",
    cc_timing: str = "before",
    clicker: str = "U1",
) -> dict[str, Any]:
    def sel(v: str) -> dict[str, Any]:
        return {
            "selected_option": {"value": v, "text": {"type": "plain_text", "text": v}}
        }

    def txt(v: str) -> dict[str, Any]:
        return {"value": v}

    state = {
        "b_amp": {"amp_email": txt(amp)},
        "b_cp": {"cp_email": txt(cp)},
        "b_seq": {"seq": sel(order)},
        "b_cc": {"cc": txt(cc)},
        "b_cc_seq": {"cc_seq": sel(cc_timing)},
    }
    return {
        "type": "view_submission",
        "user": {"id": clicker},
        "trigger_id": "trigM",
        "view": {
            "callback_id": "nda_docusign",
            "private_metadata": ref,
            "state": {"values": state},
        },
    }


def _blocks_posts(stub: StubWebClient) -> list[dict[str, Any]]:
    return [m for m in stub.messages if m.get("blocks")]


def _texts(stub: StubWebClient) -> list[str]:
    return [m.get("text", "") for m in stub.messages]


# =========================================================================== #
# 1) send_docusign → open the nda_docusign modal (no send)
# =========================================================================== #
def test_send_docusign_opens_modal(bot_session_factory) -> None:
    ref = "R1"
    _seed(
        bot_session_factory, ref, _state(signer_emails=[])
    )  # <2 route (or a generated doc)
    views: list[tuple[str, dict]] = []
    env_deps = EnvelopeDeps(
        open_view=lambda tid, view: views.append((tid, view)),
        slack_fetch=lambda att: CLEAN_DOCX,
        sender=FakeSender(),
    )
    stub = _run(
        bot_session_factory,
        _actions_body("send_docusign", send_docusign_value(ref)),
        env_deps,
    )

    assert len(views) == 1
    trigger_id, view = views[0]
    assert trigger_id == "trig"
    assert view["callback_id"] == "nda_docusign"
    assert view["private_metadata"] == ref
    # A modal open is not a thread message, and nothing is sent.
    assert stub.messages == []
    assert _envelopes(bot_session_factory) == []


def test_send_docusign_missing_state_expired(bot_session_factory) -> None:
    views: list[Any] = []
    env_deps = EnvelopeDeps(
        open_view=lambda tid, view: views.append(view), sender=FakeSender()
    )
    stub = _run(
        bot_session_factory,
        _actions_body("send_docusign", send_docusign_value("nope")),
        env_deps,
    )
    assert views == []
    assert EXPIRED_STATE_TEXT in _texts(stub)
    assert stub.messages[0]["channel"] == "C1"  # into the interaction thread


# =========================================================================== #
# 2) modal submit → golden parse + validation → SAME confirm card (still no send)
# =========================================================================== #
def test_modal_submit_parses_and_posts_confirm_card(bot_session_factory) -> None:
    ref = "R2"
    _seed(bot_session_factory, ref, _state(signer_emails=[]))
    env_deps = EnvelopeDeps(sender=FakeSender(), slack_fetch=lambda att: CLEAN_DOCX)
    stub = _run(
        bot_session_factory,
        _modal_body(
            ref,
            amp="amp@a.com",
            cp="cp@b.com",
            order="cp_first",
            cc="x@y.com, z@w.com; q@r.com",
            cc_timing="before",
        ),
        env_deps,
    )

    posts = _blocks_posts(stub)
    assert len(posts) == 1
    assert (
        posts[0]["channel"] == "C1"
    )  # the STORED origin thread (modal has no channel of its own)
    buttons = _buttons(posts[0]["blocks"])
    assert "env_confirm_send" in buttons
    assert json.loads(buttons["env_confirm_send"])["ref"] == ref

    # The collected details are merged into the durable state (golden parse).
    state = _load(bot_session_factory, ref)
    assert state is not None
    assert state["signer_emails"] == ["amp@a.com", "cp@b.com"]
    assert state["routing"] == "cp_first"
    assert state["cc_emails"] == ["x@y.com", "z@w.com", "q@r.com"]
    assert state["cc_timing"] == "before"
    # The modal is a COLLECTOR — nothing is sent here.
    assert _envelopes(bot_session_factory) == []


def test_modal_submit_invalid_email_asks_to_retry(bot_session_factory) -> None:
    ref = "R2b"
    _seed(bot_session_factory, ref, _state(signer_emails=[]))
    env_deps = EnvelopeDeps(sender=FakeSender(), slack_fetch=lambda att: CLEAN_DOCX)
    stub = _run(bot_session_factory, _modal_body(ref, amp="not-an-email"), env_deps)
    assert MODAL_INVALID_TEXT in _texts(stub)
    assert _blocks_posts(stub) == []  # no confirm card on bad input


def test_modal_submit_bad_enums_fall_back_to_defaults(bot_session_factory) -> None:
    ref = "R2c"
    _seed(bot_session_factory, ref, _state(signer_emails=[]))
    env_deps = EnvelopeDeps(sender=FakeSender(), slack_fetch=lambda att: CLEAN_DOCX)
    _run(
        bot_session_factory,
        _modal_body(ref, order="garbage", cc_timing="whenever"),
        env_deps,
    )
    state = _load(bot_session_factory, ref)
    assert state is not None
    assert state["routing"] == "all_at_once"  # unknown order → parallel default
    assert state["cc_timing"] == "after"  # unknown timing → after default


# =========================================================================== #
# 3) Confirm & send → the ONLY send path (+ requester-mapping row, PLAN §3.10)
# =========================================================================== #
def test_confirm_send_sends_and_persists_requester_mapping(bot_session_factory) -> None:
    ref = "R3"
    _seed(bot_session_factory, ref, _state(cc_emails=["cc@x.com"]))
    result = EnvelopeResult(
        envelope_id="ENV-9", status="sent", idempotency_key="IDEM1", recipients={}
    )
    sender = FakeSender(result=result)
    env_deps = EnvelopeDeps(sender=sender, slack_fetch=lambda att: CLEAN_DOCX)
    stub = _run(
        bot_session_factory,
        _actions_body("env_confirm_send", env_confirm_value(ref, "send")),
        env_deps,
    )

    # Sent exactly once, with the stored routing/signers/CC + document.
    assert len(sender.calls) == 1
    call = sender.calls[0]
    assert call["signers"] == ["amp@a.com", "cp@b.com"]
    assert call["routing"] == "amp_first"
    assert call["cc"] == ["cc@x.com"]
    assert call["cc_timing"] == "after"
    assert call["filename"] == "nda.docx"
    assert call["docx_bytes"] == CLEAN_DOCX
    assert any("Sent to DocuSign" in t for t in _texts(stub))

    # The audit + requester-mapping row (what the P4 watcher DMs from).
    envelopes = _envelopes(bot_session_factory)
    assert len(envelopes) == 1
    row = envelopes[0]
    assert (row.status, row.envelope_id, row.idempotency_key) == (
        "sent",
        "ENV-9",
        "IDEM1",
    )
    assert row.channel == "slack"
    assert row.requested_by == "U1"
    assert row.slack_channel == "C1"
    assert row.slack_thread_ts == "T1"
    assert row.routing == "amp_first"
    assert row.signer_emails == ["amp@a.com", "cp@b.com"]
    assert row.cc_emails == ["cc@x.com"]

    # State marked sent so a re-click is a no-op.
    assert _load(bot_session_factory, ref)["sent_envelope_id"] == "ENV-9"


def test_confirm_send_double_click_is_idempotent(bot_session_factory) -> None:
    ref = "R3b"
    _seed(bot_session_factory, ref, _state())
    result = EnvelopeResult(
        envelope_id="ENV-1", status="sent", idempotency_key="K1", recipients={}
    )
    sender = FakeSender(result=result)
    env_deps = EnvelopeDeps(sender=sender, slack_fetch=lambda att: CLEAN_DOCX)
    body = _actions_body("env_confirm_send", env_confirm_value(ref, "send"))

    _run(bot_session_factory, body, env_deps)
    stub2 = _run(bot_session_factory, body, env_deps)

    assert len(sender.calls) == 1  # never a second envelope
    assert ALREADY_SENT_TEXT in _texts(stub2)
    assert len(_envelopes(bot_session_factory)) == 1


def test_confirm_cancel_does_not_send(bot_session_factory) -> None:
    ref = "R4"
    _seed(bot_session_factory, ref, _state())
    sender = FakeSender(result=EnvelopeResult("E", "sent", "K", {}))
    env_deps = EnvelopeDeps(sender=sender, slack_fetch=lambda att: CLEAN_DOCX)
    stub = _run(
        bot_session_factory,
        _actions_body("env_confirm_cancel", env_confirm_value(ref, "cancel")),
        env_deps,
    )
    assert sender.calls == []
    assert CANCEL_TEXT in _texts(stub)
    assert _envelopes(bot_session_factory) == []


def test_confirm_send_missing_state_is_expired(bot_session_factory) -> None:
    sender = FakeSender(result=EnvelopeResult("E", "sent", "K", {}))
    env_deps = EnvelopeDeps(sender=sender, slack_fetch=lambda att: CLEAN_DOCX)
    stub = _run(
        bot_session_factory,
        _actions_body("env_confirm_send", env_confirm_value("gone", "send")),
        env_deps,
    )
    assert sender.calls == []
    assert EXPIRED_STATE_TEXT in _texts(stub)


def test_confirm_send_capability_off_is_friendly(bot_session_factory) -> None:
    ref = "R5"
    _seed(bot_session_factory, ref, _state())
    # No injected sender → the handler builds from settings; a docusign-less Settings => DOCUSIGN disabled.
    env_deps = EnvelopeDeps(slack_fetch=lambda att: CLEAN_DOCX)
    settings = Settings(_env_file=None, nda_bot_from_email="nda-bot@example.com")  # type: ignore[call-arg]
    stub = _run(
        bot_session_factory,
        _actions_body("env_confirm_send", env_confirm_value(ref, "send")),
        env_deps,
        settings=settings,
    )
    assert DOCUSIGN_UNAVAILABLE_TEXT in _texts(stub)
    assert _envelopes(bot_session_factory) == []  # capability off => nothing recorded


def test_confirm_send_failure_persists_failed_attempt(bot_session_factory) -> None:
    ref = "R6"
    _seed(bot_session_factory, ref, _state())
    sender = FakeSender(
        exc=DocuSignTerminalError(
            "bad", status_code=400, error_code="INVALID_EMAIL_ADDRESS"
        )
    )
    env_deps = EnvelopeDeps(sender=sender, slack_fetch=lambda att: CLEAN_DOCX)
    stub = _run(
        bot_session_factory,
        _actions_body("env_confirm_send", env_confirm_value(ref, "send")),
        env_deps,
    )

    assert any("DocuSign didn't accept" in t for t in _texts(stub))
    envelopes = _envelopes(bot_session_factory)
    assert len(envelopes) == 1
    row = envelopes[0]
    assert row.status == "failed"
    assert row.envelope_id is None
    assert (
        row.idempotency_key
    )  # computed for the failed attempt so the audit still keys the doc+recipients
    assert row.requested_by == "U1"  # requester mapping recorded even on failure


# =========================================================================== #
# 4) Thread-doc confirm chain — env_use_doc / decline_doc
# =========================================================================== #
def test_env_use_doc_fetches_and_posts_confirm_for_two_signers(
    bot_session_factory,
) -> None:
    ref = "R7"
    _seed(
        bot_session_factory,
        ref,
        _state(
            signer_emails=["a@a.com", "b@b.com"], doc_b64=None, slack_file_id="Fthread"
        ),
    )
    # remove the None doc_b64 the helper set (thread doc carries only a ref)
    with bot_session_factory() as s, s.begin():
        row = s.query(BotCorrelation).filter(BotCorrelation.key == ref).one()
        payload = dict(row.payload_json)
        payload.pop("doc_b64", None)
        row.payload_json = payload

    env_deps = EnvelopeDeps(slack_fetch=lambda att: CLEAN_DOCX, sender=FakeSender())
    stub = _run(
        bot_session_factory,
        _actions_body("env_use_doc", env_use_doc_value(ref)),
        env_deps,
    )

    posts = _blocks_posts(stub)
    assert len(posts) == 1
    assert "env_confirm_send" in _buttons(posts[0]["blocks"])
    # The fetched bytes are inlined so confirm doesn't re-fetch a possibly-deleted file.
    assert base64.b64decode(_load(bot_session_factory, ref)["doc_b64"]) == CLEAN_DOCX
    assert _envelopes(bot_session_factory) == []


def test_env_use_doc_under_two_signers_posts_signer_button(bot_session_factory) -> None:
    ref = "R7b"
    _seed(
        bot_session_factory,
        ref,
        _state(signer_emails=[], doc_b64=None, slack_file_id="Fthread"),
    )
    with bot_session_factory() as s, s.begin():
        row = s.query(BotCorrelation).filter(BotCorrelation.key == ref).one()
        payload = dict(row.payload_json)
        payload.pop("doc_b64", None)
        row.payload_json = payload

    env_deps = EnvelopeDeps(slack_fetch=lambda att: CLEAN_DOCX, sender=FakeSender())
    stub = _run(
        bot_session_factory,
        _actions_body("env_use_doc", env_use_doc_value(ref)),
        env_deps,
    )
    assert "send_docusign" in _buttons(_blocks_posts(stub)[0]["blocks"])


def test_env_use_doc_tokenised_is_refused(bot_session_factory) -> None:
    ref = "R7c"
    _seed(
        bot_session_factory,
        ref,
        _state(signer_emails=["a@a.com", "b@b.com"], slack_file_id="Fthread"),
    )
    with bot_session_factory() as s, s.begin():
        row = s.query(BotCorrelation).filter(BotCorrelation.key == ref).one()
        payload = dict(row.payload_json)
        payload.pop("doc_b64", None)
        row.payload_json = payload

    env_deps = EnvelopeDeps(slack_fetch=lambda att: TOKENISED_DOCX, sender=FakeSender())
    stub = _run(
        bot_session_factory,
        _actions_body("env_use_doc", env_use_doc_value(ref)),
        env_deps,
    )
    assert any("unfilled placeholders" in t for t in _texts(stub))
    assert _blocks_posts(stub) == []  # no confirm card for a tokenised doc


def test_decline_doc_replies_attach_a_file(bot_session_factory) -> None:
    env_deps = EnvelopeDeps(sender=FakeSender(), slack_fetch=lambda att: CLEAN_DOCX)
    stub = _run(
        bot_session_factory,
        _actions_body("decline_doc", decline_doc_value("R8")),
        env_deps,
    )
    assert DECLINE_TEXT in _texts(stub)
    assert _envelopes(bot_session_factory) == []
