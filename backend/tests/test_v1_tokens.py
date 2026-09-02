"""HTTP tests for the in-Word tokenizer plane (SERVICE-authed, ``app.api.routes_tokens_v1``):

  * ``GET  /v1/tokens`` — the registry palette + exact ``{{name}}`` placeholder text.
  * ``POST /v1/support_task/template-draft`` — land a tokenised .docx as a NEW DRAFT template_version.

Both reuse the same authorization ``/v1/support_task/generate-nda`` enforces (``engine_principal`` +
the >=1-engine-entitlement gate). The engine is unconfigured in tests, so ``engine_principal`` binds
the open ``svc:local`` dev principal (full entitlements) — except where a test CONFIGURES the engine
(to prove the 401 no-key path) or OVERRIDES the principal (to prove the 403 no-entitlement path). No
network: the registry + templates are seeded through the throwaway-DB ``db`` session.
"""

from __future__ import annotations

import io

import pytest
from conftest_admin import seed_catalog, template_id
from docx import Document

from app.models_v2 import TemplateVersion

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_bytes(text: str = "{{counterparty_name}}") -> bytes:
    d = Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _draft_files(
    name: str = "tokenised.docx", body: bytes | None = None, mime: str = _DOCX_MIME
):
    return {"file": (name, io.BytesIO(_docx_bytes() if body is None else body), mime)}


@pytest.fixture
def configured_engine(client, monkeypatch):
    """Configure the engine so an unauthenticated /v1 call is rejected 401 (not bound to the open
    svc:local dev principal). Mirrors ``tests/test_engine_review_gate.configured_engine``."""
    from app.config import settings

    monkeypatch.setattr(settings, "engine_api_key", "test-engine-key", raising=False)
    monkeypatch.setattr(settings, "engine_service_keys", "", raising=False)
    return client


def _override_noent(app):
    """Override ``engine_principal`` on the live app with a SERVICE principal holding NO entitlements."""
    from app.auth.principal import ResolvedPrincipal, engine_principal
    from app.schemas import DEFAULT_ORG_ID

    app.dependency_overrides[engine_principal] = lambda: ResolvedPrincipal(
        principal_type="service",
        principal_id="noent",
        org_id=DEFAULT_ORG_ID,
        entitlements=frozenset(),
    )


# --------------------------------------------------------------------------- #
# GET /v1/tokens
# --------------------------------------------------------------------------- #
def test_tokens_returns_seeded_registry_with_labels_and_placeholders(client, db):
    seed_catalog(db)
    resp = client.get("/v1/tokens")
    assert resp.status_code == 200, resp.text
    tokens = resp.json()["tokens"]
    assert isinstance(tokens, list) and len(tokens) == 16  # the 16 seed tokens
    by_name = {t["name"]: t for t in tokens}
    cp = by_name["counterparty_name"]
    # The exact {{name}} form the add-in splices in, plus the human label (never raw snake_case).
    assert cp["placeholder"] == "{{counterparty_name}}"
    assert cp["label"] == "Counterparty Name"
    # Every documented field is present on each token.
    assert set(cp) == {
        "name",
        "label",
        "help_text",
        "data_type",
        "party",
        "scope_code",
        "placeholder",
    }
    assert all(t["placeholder"] == "{{" + t["name"] + "}}" for t in tokens)


def test_tokens_requires_api_key_when_configured(configured_engine):
    resp = configured_engine.get("/v1/tokens")  # no X-API-Key on a configured engine
    assert resp.status_code == 401
    assert resp.json()["error"]["code"] == "unauthorized"


def test_tokens_requires_an_engine_entitlement(client, app, db):
    seed_catalog(db)
    _override_noent(app)
    try:
        resp = client.get("/v1/tokens")
        assert resp.status_code == 403
        assert resp.json()["error"]["code"] == "not_entitled"
    finally:
        from app.auth.principal import engine_principal

        app.dependency_overrides.pop(engine_principal, None)


# --------------------------------------------------------------------------- #
# POST /v1/support_task/template-draft
# --------------------------------------------------------------------------- #
def test_template_draft_stores_a_draft_version(client, db):
    seed_catalog(db)
    tid = template_id(db, "US", "ServiceProvider", "NotApplicable")
    resp = client.post(
        "/v1/support_task/template-draft",
        files=_draft_files(),
        data={
            "jurisdiction": "us",  # lenient codes normalise like generate-nda
            "counterparty_type": "service_provider",
            "mutuality": "",
        },
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["template"] == "US Service Provider NDA"
    assert body["variant"] == "tokenised"
    assert body["version_no"] == 2  # seed v1 exists; this is the next draft

    version = db.get(TemplateVersion, body["version_id"])
    assert version is not None
    assert version.template_id == tid
    assert version.is_current is False  # a DRAFT — never published here
    assert version.created_by == "svc:local"  # attributed to the SERVICE principal


def test_template_draft_does_not_flip_is_current(client, db):
    seed_catalog(db)
    tid = template_id(db, "US", "ServiceProvider", "NotApplicable")
    # The seeded tokenised v1 is the current one (blob NULL).
    seeded_current = (
        db.query(TemplateVersion)
        .filter(
            TemplateVersion.template_id == tid,
            TemplateVersion.variant_code == "tokenised",
            TemplateVersion.is_current.is_(True),
        )
        .one()
    )
    resp = client.post(
        "/v1/support_task/template-draft",
        files=_draft_files(),
        data={"jurisdiction": "US", "counterparty_type": "ServiceProvider"},
    )
    assert resp.status_code == 201, resp.text
    new_id = resp.json()["version_id"]

    db.expire_all()
    # The seeded current is STILL current; the new draft is not; exactly one current remains.
    assert db.get(TemplateVersion, seeded_current.id).is_current is True
    assert db.get(TemplateVersion, new_id).is_current is False
    currents = (
        db.query(TemplateVersion)
        .filter(
            TemplateVersion.template_id == tid,
            TemplateVersion.variant_code == "tokenised",
            TemplateVersion.is_current.is_(True),
        )
        .all()
    )
    assert [c.id for c in currents] == [seeded_current.id]


def test_template_draft_rejects_unknown_jurisdiction(client, db):
    seed_catalog(db)
    resp = client.post(
        "/v1/support_task/template-draft",
        files=_draft_files(),
        data={"jurisdiction": "XX", "counterparty_type": "ServiceProvider"},
    )
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "bad_request"


def test_template_draft_rejects_unknown_combo_no_template(client, db):
    # Valid codes that normalise cleanly, but no template seeded → typed 404.
    resp = client.post(
        "/v1/support_task/template-draft",
        files=_draft_files(),
        data={"jurisdiction": "US", "counterparty_type": "ServiceProvider"},
    )
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "template_not_found"


def test_template_draft_rejects_non_docx(client, db):
    seed_catalog(db)
    resp = client.post(
        "/v1/support_task/template-draft",
        files={"file": ("nda.txt", io.BytesIO(b"not a docx"), "text/plain")},
        data={"jurisdiction": "US", "counterparty_type": "ServiceProvider"},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "not_docx"


def test_template_draft_rejects_docx_named_non_docx(client, db):
    # A .docx-named file whose bytes are not a readable OOXML doc → the studio bad-docx refusal.
    seed_catalog(db)
    resp = client.post(
        "/v1/support_task/template-draft",
        files={"file": ("nda.docx", io.BytesIO(b"not really a docx"), _DOCX_MIME)},
        data={"jurisdiction": "US", "counterparty_type": "ServiceProvider"},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "studio_bad_docx"


def test_template_draft_rejects_empty_file(client, db):
    seed_catalog(db)
    resp = client.post(
        "/v1/support_task/template-draft",
        files={"file": ("nda.docx", io.BytesIO(b""), _DOCX_MIME)},
        data={"jurisdiction": "US", "counterparty_type": "ServiceProvider"},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "empty_file"


def test_template_draft_rejects_oversize(client, db, monkeypatch):
    seed_catalog(db)
    from app.config import settings

    # Drive max_upload_bytes (a computed property of max_upload_mb) to 0 so any content trips the cap.
    monkeypatch.setattr(settings, "max_upload_mb", 0, raising=False)
    resp = client.post(
        "/v1/support_task/template-draft",
        files=_draft_files(),
        data={"jurisdiction": "US", "counterparty_type": "ServiceProvider"},
    )
    assert resp.status_code == 413
    assert resp.json()["error"]["code"] == "request_too_large"


def test_template_draft_requires_api_key_when_configured(configured_engine):
    resp = configured_engine.post(
        "/v1/support_task/template-draft",
        files=_draft_files(),
        data={"jurisdiction": "US", "counterparty_type": "ServiceProvider"},
    )
    assert resp.status_code == 401
    assert resp.json()["error"]["code"] == "unauthorized"


def test_template_draft_requires_an_engine_entitlement(client, app, db):
    seed_catalog(db)
    _override_noent(app)
    try:
        resp = client.post(
            "/v1/support_task/template-draft",
            files=_draft_files(),
            data={"jurisdiction": "US", "counterparty_type": "ServiceProvider"},
        )
        assert resp.status_code == 403
        assert resp.json()["error"]["code"] == "not_entitled"
    finally:
        from app.auth.principal import engine_principal

        app.dependency_overrides.pop(engine_principal, None)
