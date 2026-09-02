"""Shared builders for the admin-shell + template-studio page tests.

Plain helpers importable as ``from conftest_admin import ...`` (pytest's prepend import mode puts
``tests/`` on ``sys.path`` — the ``conftest_studio`` convention). No network, no fixtures on disk;
documents are built with python-docx in memory and persisted through the throwaway-DB ``db`` session.

The app fixture does NOT run the lifespan (so the real seed never fires), so these helpers seed the 8
templates + 16 tokens (``seed_catalog``) and build draft ``template_version`` rows with content-
addressed blobs, exactly as the studio upload route does.
"""

from __future__ import annotations

import hashlib
import uuid
from io import BytesIO
from typing import Any

from docx import Document
from sqlalchemy import func, select


# --------------------------------------------------------------------------- #
# Document builders
# --------------------------------------------------------------------------- #
def _to_bytes(doc: Any) -> bytes:
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def build_docx(*paragraphs: str) -> bytes:
    """A .docx with one body paragraph per string (each a single run — predictable offsets)."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    return _to_bytes(doc)


def source_docx() -> bytes:
    """A source document a template author would tokenise: a hardcoded company name to highlight→click
    and a typed ``[EFFECTIVE DATE]`` placeholder for the find-and-map assistant."""
    return build_docx(
        "This Agreement is between ACME CORPORATION and the Recipient.",
        "Effective as of [EFFECTIVE DATE].",
    )


def all_tokens_docx() -> bytes:
    """A .docx already containing every seed token placeholder — so ``missing_required`` is empty for
    any variant/scope (used to exercise the publish success path)."""
    from app.seed_catalog import TOKENS

    return build_docx(*[f"{{{{{t['name']}}}}}" for t in TOKENS])


def bare_docx() -> bytes:
    """A .docx with no tokens at all — ``missing_required`` is non-empty for a tokenised variant."""
    return build_docx("A plain paragraph with no placeholders at all.")


def tokens_docx(*names: str) -> bytes:
    """A .docx whose paragraphs each hold one ``{{name}}`` placeholder."""
    return build_docx(*[f"{{{{{n}}}}}" for n in names])


# --------------------------------------------------------------------------- #
# Seeding
# --------------------------------------------------------------------------- #
def seed_catalog(db: Any) -> None:
    """Seed the 8 templates (+ empty/tokenised v1) + 16 tokens + token_template mapping (idempotent)."""
    from app.schemas import DEFAULT_ORG_ID
    from app.seed_catalog import seed_templates_tokens

    seed_templates_tokens(db, DEFAULT_ORG_ID)
    db.commit()


def template_id(
    db: Any,
    jurisdiction: str = "US",
    counterparty_type: str = "ServiceProvider",
    mutuality: str = "NotApplicable",
) -> str:
    from app.models_v2 import Template

    return (
        db.execute(
            select(Template).where(
                Template.jurisdiction_code == jurisdiction,
                Template.counterparty_type_code == counterparty_type,
                Template.mutuality_code == mutuality,
            )
        )
        .scalar_one()
        .id
    )


def _store_blob(db: Any, data: bytes) -> str:
    from app.models_v2 import DocumentBlob
    from app.support_task.generator import DOCX_MIME

    sha = hashlib.sha256(data).hexdigest()
    blob = db.execute(
        select(DocumentBlob).where(DocumentBlob.sha256 == sha)
    ).scalar_one_or_none()
    if blob is None:
        blob = DocumentBlob(
            sha256=sha, byte_size=len(data), mime_type=DOCX_MIME, bytes=data
        )
        db.add(blob)
        db.flush()
    return blob.id


def add_version(
    db: Any,
    template_id: str,
    variant: str,
    docx_bytes: bytes,
    *,
    is_current: bool = False,
    version_no: int | None = None,
) -> str:
    """Persist a new ``template_version`` (content-addressed blob) for a (template, variant) and return
    its id. ``version_no`` defaults to the next available number for the slot."""
    from app.models_v2 import TemplateVersion

    if version_no is None:
        version_no = (
            int(
                db.execute(
                    select(func.max(TemplateVersion.version_no)).where(
                        TemplateVersion.template_id == template_id,
                        TemplateVersion.variant_code == variant,
                    )
                ).scalar()
                or 0
            )
            + 1
        )
    version = TemplateVersion(
        id=uuid.uuid4().hex,
        template_id=template_id,
        variant_code=variant,
        version_no=version_no,
        blob_id=_store_blob(db, docx_bytes),
        is_current=is_current,
    )
    db.add(version)
    db.commit()
    return version.id


def seed_draft(
    db: Any,
    docx_bytes: bytes,
    *,
    variant: str = "tokenised",
    jurisdiction: str = "US",
    counterparty_type: str = "ServiceProvider",
    mutuality: str = "NotApplicable",
) -> tuple[str, str]:
    """Seed the catalog + a draft version of ``docx_bytes`` for one template slot. Returns
    ``(template_id, version_id)``."""
    seed_catalog(db)
    tid = template_id(db, jurisdiction, counterparty_type, mutuality)
    vid = add_version(db, tid, variant, docx_bytes)
    return tid, vid


def span_of(docx_bytes: bytes, locator: str, needle: str) -> tuple[int, int, str]:
    """The ``(start, end, content_hash)`` for ``needle`` inside the paragraph at ``locator`` — computed
    from the same ``extract_view`` the studio page renders from, so a tokenize POST built off it is
    byte-accurate."""
    from app.studio.docview import extract_view

    view = extract_view(docx_bytes)
    seg = view.find(locator)
    assert seg is not None, f"no segment {locator!r}"
    start = seg.text.index(needle)
    return start, start + len(needle), view.content_hash
